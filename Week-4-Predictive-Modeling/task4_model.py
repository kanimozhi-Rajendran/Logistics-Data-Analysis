
from pathlib import Path
import warnings

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt

from sklearn.model_selection import train_test_split, KFold, cross_val_score
from sklearn.compose import ColumnTransformer
from sklearn.preprocessing import OneHotEncoder
from sklearn.impute import SimpleImputer
from sklearn.pipeline import Pipeline
from sklearn.linear_model import LinearRegression
from sklearn.tree import DecisionTreeRegressor
from sklearn.metrics import mean_absolute_error, mean_squared_error, r2_score

from docx import Document
from docx.shared import Inches

warnings.filterwarnings("ignore")

# Project folders
PROJECT = Path(__file__).resolve().parent.parent
DATA = PROJECT / "data" / "logistics_data.csv"
OUTPUT = PROJECT / "outputs" / "week4"
CHARTS = OUTPUT / "charts"

OUTPUT.mkdir(parents=True, exist_ok=True)
CHARTS.mkdir(parents=True, exist_ok=True)

TARGET = "Delivery_Days"

NUMERIC = [
    "Distance_km",
    "Expected_Days",
    "Shipping_Cost",
    "Order_Quantity"
]

CATEGORICAL = [
    "Traffic_Level",
    "Weather",
    "Vehicle_Type"
]

FEATURES = NUMERIC + CATEGORICAL


def create_model(model):
    numeric_steps = Pipeline([
        ("imputer", SimpleImputer(strategy="median"))
    ])

    category_steps = Pipeline([
        ("imputer", SimpleImputer(strategy="most_frequent")),
        ("encoder", OneHotEncoder(handle_unknown="ignore"))
    ])

    preparation = ColumnTransformer([
        ("numeric", numeric_steps, NUMERIC),
        ("categorical", category_steps, CATEGORICAL)
    ])

    return Pipeline([
        ("preparation", preparation),
        ("model", model)
    ])


def main():
    # 1. Load dataset
    if not DATA.exists():
        raise FileNotFoundError(f"Dataset not found: {DATA}")

    df = pd.read_csv(DATA)
    df.columns = df.columns.str.strip()

    required = FEATURES + [TARGET]
    missing = [c for c in required if c not in df.columns]

    if missing:
        raise ValueError(f"Missing dataset columns: {missing}")

    for col in NUMERIC + [TARGET]:
        df[col] = pd.to_numeric(df[col], errors="coerce")

    for col in CATEGORICAL:
        df[col] = df[col].astype("string").str.strip()

    df = df.dropna(subset=[TARGET]).copy()

    X = df[FEATURES]
    y = df[TARGET]

    if len(df) < 8:
        raise ValueError("At least 8 usable records are needed.")

    print("Dataset shape:", df.shape)
    print("Target: Delivery_Days")

    # 2. Split data into training and testing
    X_train, X_test, y_train, y_test = train_test_split(
        X, y, test_size=0.25, random_state=42
    )

    # 3. Compare two models
    models = {
        "Linear Regression": create_model(LinearRegression()),
        "Decision Tree": create_model(
            DecisionTreeRegressor(
                max_depth=2,
                min_samples_leaf=2,
                random_state=42
            )
        )
    }

    # Cross-validation uses training data only
    cv = KFold(n_splits=3, shuffle=True, random_state=42)
    scores = []

    for name, model in models.items():
        fold_scores = cross_val_score(
            model,
            X_train,
            y_train,
            cv=cv,
            scoring="neg_mean_absolute_error"
        )

        scores.append({
            "Model": name,
            "CV_MAE": -fold_scores.mean(),
            "CV_MAE_Std": fold_scores.std()
        })

    comparison = pd.DataFrame(scores).sort_values("CV_MAE")
    comparison.to_csv(
        OUTPUT / "cross_validation_results.csv", index=False
    )

    best_name = comparison.iloc[0]["Model"]
    best_model = models[best_name]
    best_model.fit(X_train, y_train)

    # 4. Test the selected model
    predictions = best_model.predict(X_test)

    mae = mean_absolute_error(y_test, predictions)
    rmse = np.sqrt(mean_squared_error(y_test, predictions))
    r2 = r2_score(y_test, predictions)

    results = pd.DataFrame({
        "Actual_Delivery_Days": y_test.to_numpy(),
        "Predicted_Delivery_Days": np.round(predictions, 2),
        "Absolute_Error": np.round(
            np.abs(y_test.to_numpy() - predictions), 2
        )
    })

    results.to_csv(OUTPUT / "test_predictions.csv", index=False)

    metrics = pd.DataFrame([{
        "Selected_Model": best_name,
        "Test_MAE": mae,
        "Test_RMSE": rmse,
        "Test_R2": r2,
        "Training_Rows": len(X_train),
        "Testing_Rows": len(X_test)
    }])

    metrics.to_csv(OUTPUT / "model_evaluation.csv", index=False)

    # 5. Chart: model comparison
    plt.figure(figsize=(7, 5))
    plt.bar(comparison["Model"], comparison["CV_MAE"],
            edgecolor="black")
    plt.title("Model Comparison Using Cross-Validation MAE")
    plt.ylabel("Mean Absolute Error (days)")
    plt.xlabel("Model")
    plt.tight_layout()
    plt.savefig(CHARTS / "model_comparison.png", dpi=150)
    plt.close()

    # 6. Chart: actual versus predicted
    plt.figure(figsize=(7, 5))
    plt.scatter(y_test, predictions)

    low = min(float(y_test.min()), float(np.min(predictions)))
    high = max(float(y_test.max()), float(np.max(predictions)))

    plt.plot([low, high], [low, high], linestyle="--")
    plt.title("Actual vs Predicted Delivery Time")
    plt.xlabel("Actual Delivery Days")
    plt.ylabel("Predicted Delivery Days")
    plt.tight_layout()
    plt.savefig(CHARTS / "actual_vs_predicted.png", dpi=150)
    plt.close()

    # 7. Create Word report
    doc = Document()

    doc.add_heading("LOGISTICS DATA ANALYSIS", 0)
    doc.add_heading(
        "Week 4: Predictive Modeling and Optimization", level=1
    )
    doc.add_paragraph("Prepared by: Kanimozhi R")
    doc.add_paragraph(
        "Dataset: Hypothetical logistics dataset"
    )

    doc.add_heading("1. Problem Definition", level=1)
    doc.add_paragraph(
        "The objective is to predict delivery time in days using "
        "logistics information. Accurate estimates can help teams "
        "plan deliveries, identify potentially delayed orders, and "
        "improve operational efficiency."
    )

    doc.add_heading("2. Dataset and Features", level=1)
    doc.add_paragraph(
        f"The dataset contains {len(df)} usable orders. "
        "Delivery_Days is the target variable. Input features include "
        "distance, expected delivery days, shipping cost, order quantity, "
        "traffic level, weather, and vehicle type. Missing input values "
        "are handled through median or most-frequent-value imputation. "
        "Categorical variables are converted using one-hot encoding."
    )

    doc.add_heading("3. Model Selection", level=1)
    doc.add_paragraph(
        "Linear Regression is used as a simple baseline model. "
        "A Decision Tree Regressor is also tested because it can model "
        "non-linear patterns. The tree depth and minimum leaf size "
        "are limited to reduce overfitting. Three-fold cross-validation "
        "compares models using Mean Absolute Error (MAE)."
    )

    doc.add_heading("4. Training and Validation", level=1)
    for step in [
        "Separate Delivery_Days from the input features.",
        "Split the data into 75% training and 25% testing.",
        "Impute missing values and encode categories inside a pipeline.",
        "Compare the two models using three-fold cross-validation.",
        "Evaluate the selected model on the held-out test data."
    ]:
        doc.add_paragraph(step, style="List Number")

    doc.add_heading("5. Model Performance", level=1)

    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Metric"
    table.rows[0].cells[1].text = "Result"

    metric_rows = [
        ("Selected model", best_name),
        ("Training records", str(len(X_train))),
        ("Testing records", str(len(X_test))),
        ("Cross-validation MAE (days)",
         f"{comparison.iloc[0]['CV_MAE']:.3f}"),
        ("Test MAE (days)", f"{mae:.3f}"),
        ("Test RMSE (days)", f"{rmse:.3f}"),
        ("Test R-squared", f"{r2:.3f}")
    ]

    for label, value in metric_rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value

    doc.add_paragraph(
        "MAE measures the average absolute prediction error. RMSE "
        "penalizes larger errors more strongly. Lower MAE and RMSE "
        "generally indicate better predictions. R-squared compares "
        "model performance with a mean-based baseline; it can be "
        "negative on test data. Since this dataset is very small, "
        "the scores are only an initial demonstration."
    )

    doc.add_heading("6. Visualization and Interpretation", level=1)

    doc.add_heading("6.1 Model Comparison", level=2)
    doc.add_picture(
        str(CHARTS / "model_comparison.png"), width=Inches(5.7)
    )
    doc.add_paragraph(
        "This chart compares cross-validation MAE. A lower bar means "
        "the model had a lower average absolute error across the "
        "training folds."
    )

    doc.add_heading("6.2 Actual vs Predicted Delivery Time", level=2)
    doc.add_picture(
        str(CHARTS / "actual_vs_predicted.png"), width=Inches(5.7)
    )
    doc.add_paragraph(
        "Points closer to the diagonal line indicate more accurate "
        "predictions. Larger gaps show orders with larger prediction "
        "errors."
    )

    doc.add_heading("7. Logistics Optimization Strategies", level=1)

    recommendations = [
        "Review orders predicted to exceed their expected delivery window.",
        "Check traffic, weather, route distance, and dispatch time before responding to delay risks.",
        "Group deliveries to nearby destinations when capacity and promised delivery windows allow.",
        "Compare vehicle costs for similar routes before reallocating vehicles.",
        "Track actual delivery times against predictions and monitor errors regularly.",
        "Collect more historical records before using the model for operational decisions."
    ]

    for item in recommendations:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("8. Limitations", level=1)
    doc.add_paragraph(
        "The dataset has only 20 illustrative records, so the test set "
        "is very small and the results may change with a different split. "
        "Cross-validation helps compare candidate models but does not "
        "replace testing on a larger real-world dataset. Predictions "
        "should not be treated as guaranteed delivery times."
    )

    doc.add_heading("9. Conclusion", level=1)
    doc.add_paragraph(
        f"The experiment compared Linear Regression and Decision Tree "
        f"models to predict delivery time. {best_name} was selected "
        "because it had the lower mean cross-validation MAE. The "
        "evaluation and charts demonstrate a basic predictive workflow "
        "that can support logistics planning when combined with reliable "
        "data and operational judgment."
    )

    report = OUTPUT / "Week4_Predictive_Modeling_Report.docx"
    doc.save(report)

    # Save a short text summary
    summary = (
        "WEEK 4 PREDICTIVE MODELING SUMMARY\n"
        f"Records used: {len(df)}\n"
        f"Selected model: {best_name}\n"
        f"Cross-validation MAE: {comparison.iloc[0]['CV_MAE']:.3f} days\n"
        f"Test MAE: {mae:.3f} days\n"
        f"Test RMSE: {rmse:.3f} days\n"
        f"Test R-squared: {r2:.3f}\n"
        f"Word report: {report}\n"
        f"Charts folder: {CHARTS}\n"
        "Note: Small illustrative dataset; results are educational."
    )

    (OUTPUT / "model_summary.txt").write_text(
        summary, encoding="utf-8"
    )

    print("\n--- WEEK 4 MODELING COMPLETED ---")
    print(summary)


if __name__ == "__main__":
    main()
