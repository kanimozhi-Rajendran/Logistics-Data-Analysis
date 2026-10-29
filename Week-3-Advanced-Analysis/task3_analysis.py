
from pathlib import Path
import pandas as pd
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches

# Project folders
PROJECT = Path(__file__).resolve().parent.parent
DATA = PROJECT / "data" / "logistics_data.csv"
OUTPUT = PROJECT / "outputs" / "week3"
CHARTS = OUTPUT / "charts"

OUTPUT.mkdir(parents=True, exist_ok=True)
CHARTS.mkdir(parents=True, exist_ok=True)


def main():
    # 1. Load and prepare the dataset
    df = pd.read_csv(DATA)
    df.columns = df.columns.str.strip()

    numeric_cols = [
        "Distance_km", "Shipping_Cost", "Delivery_Days",
        "Expected_Days", "Order_Quantity"
    ]

    for col in numeric_cols:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")

    print("Dataset shape:", df.shape)
    print("\nFirst five rows:")
    print(df.head())

    # 2. Descriptive statistics: mean, median, spread
    numbers = df.select_dtypes(include="number")
    stats = numbers.describe().T
    stats["median"] = numbers.median()
    stats.to_csv(OUTPUT / "descriptive_statistics.csv")

    print("\nDescriptive statistics:")
    print(stats)

    # 3. Correlation analysis
    correlation = numbers.corr()
    correlation.to_csv(OUTPUT / "correlation_matrix.csv")

    # Chart 1: Delivery time distribution
    if "Delivery_Days" in df:
        plt.figure(figsize=(8, 5))
        plt.hist(df["Delivery_Days"].dropna(), bins="auto",
                 edgecolor="black")
        plt.title("Distribution of Delivery Time")
        plt.xlabel("Delivery Days")
        plt.ylabel("Number of Orders")
        plt.tight_layout()
        plt.savefig(CHARTS / "delivery_distribution.png", dpi=150)
        plt.close()

    # Chart 2: Average delivery time by traffic
    if {"Traffic_Level", "Delivery_Days"}.issubset(df.columns):
        avg_delivery = df.groupby("Traffic_Level")[
            "Delivery_Days"
        ].mean()

        plt.figure(figsize=(8, 5))
        avg_delivery.plot(kind="bar", edgecolor="black")
        plt.title("Average Delivery Time by Traffic Level")
        plt.xlabel("Traffic Level")
        plt.ylabel("Average Delivery Days")
        plt.xticks(rotation=0)
        plt.tight_layout()
        plt.savefig(CHARTS / "delivery_by_traffic.png", dpi=150)
        plt.close()

    # Chart 3: Average shipping cost by vehicle type
    if {"Vehicle_Type", "Shipping_Cost"}.issubset(df.columns):
        avg_cost = df.groupby("Vehicle_Type")[
            "Shipping_Cost"
        ].mean()

        plt.figure(figsize=(8, 5))
        avg_cost.plot(kind="bar", edgecolor="black")
        plt.title("Average Shipping Cost by Vehicle Type")
        plt.xlabel("Vehicle Type")
        plt.ylabel("Average Shipping Cost")
        plt.xticks(rotation=20)
        plt.tight_layout()
        plt.savefig(CHARTS / "cost_by_vehicle.png", dpi=150)
        plt.close()

    # Chart 4: Distance and shipping cost relationship
    if {"Distance_km", "Shipping_Cost"}.issubset(df.columns):
        plt.figure(figsize=(8, 5))
        plt.scatter(df["Distance_km"], df["Shipping_Cost"])
        plt.title("Distance vs Shipping Cost")
        plt.xlabel("Distance (km)")
        plt.ylabel("Shipping Cost")
        plt.tight_layout()
        plt.savefig(CHARTS / "distance_vs_cost.png", dpi=150)
        plt.close()

    # Chart 5: Correlation heatmap
    if not correlation.empty:
        fig, ax = plt.subplots(figsize=(8, 6))
        image = ax.imshow(correlation, vmin=-1, vmax=1)
        ax.set_xticks(range(len(correlation.columns)))
        ax.set_yticks(range(len(correlation.columns)))
        ax.set_xticklabels(
            correlation.columns, rotation=45, ha="right"
        )
        ax.set_yticklabels(correlation.columns)
        ax.set_title("Correlation Heatmap")
        fig.colorbar(image, ax=ax, label="Correlation")
        fig.tight_layout()
        fig.savefig(CHARTS / "correlation_heatmap.png", dpi=150)
        plt.close(fig)

    # 4. Calculate summary metrics
    total_orders = len(df)
    average_cost = df["Shipping_Cost"].mean()
    average_delivery = df["Delivery_Days"].mean()
    average_distance = df["Distance_km"].mean()

    on_time = None
    delayed = None
    on_time_rate = None

    if {"Delivery_Days", "Expected_Days"}.issubset(df.columns):
        valid = df[["Delivery_Days", "Expected_Days"]].dropna()
        on_time = int(
            (valid["Delivery_Days"] <= valid["Expected_Days"]).sum()
        )
        delayed = len(valid) - on_time
        if len(valid):
            on_time_rate = on_time / len(valid) * 100

    # 5. Generate Word report with charts
    doc = Document()
    doc.add_heading("LOGISTICS DATA ANALYSIS", 0)
    doc.add_heading(
        "Week 3: Advanced Data Analysis and Visualization", level=1
    )
    doc.add_paragraph("Prepared by: Kanimozhi R")
    doc.add_paragraph("Dataset: Hypothetical logistics sample")

    doc.add_heading("1. Introduction", level=1)
    doc.add_paragraph(
        "This project uses Python to explore logistics data and visualize "
        "delivery performance, shipping cost, transportation distance, "
        "traffic, and vehicle types. The dataset is an illustrative sample."
    )

    doc.add_heading("2. Dataset and Methodology", level=1)
    doc.add_paragraph(
        f"The dataset contains {len(df)} orders and {len(df.columns)} "
        "columns. Pandas is used to load and analyze the data. Mean, median, "
        "descriptive statistics, and correlation are calculated. Matplotlib "
        "is used to create charts."
    )

    doc.add_heading("3. Key Results", level=1)
    metrics = [
        ("Total orders", str(total_orders)),
        ("Average shipping cost", f"{average_cost:.2f}"),
        ("Average delivery time (days)", f"{average_delivery:.2f}"),
        ("Average distance (km)", f"{average_distance:.2f}")
    ]
    if on_time_rate is not None:
        metrics.extend([
            ("Orders delivered on time", str(on_time)),
            ("Delayed orders", str(delayed)),
            ("On-time delivery rate", f"{on_time_rate:.1f}%")
        ])

    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Metric"
    table.rows[0].cells[1].text = "Result"
    for name, value in metrics:
        cells = table.add_row().cells
        cells[0].text = name
        cells[1].text = value

    chart_details = [
        ("delivery_distribution.png", "4.1 Delivery Time Distribution",
         "The histogram shows how delivery times are distributed across "
         "orders and helps identify common delivery durations."),
        ("delivery_by_traffic.png", "4.2 Delivery Time by Traffic",
         "The bar chart compares average delivery days for traffic levels. "
         "Higher values may indicate a need to investigate congestion."),
        ("cost_by_vehicle.png", "4.3 Shipping Cost by Vehicle",
         "The bar chart compares average shipping costs across vehicle "
         "types and supports further cost-efficiency investigation."),
        ("distance_vs_cost.png", "4.4 Distance and Shipping Cost",
         "The scatter plot helps examine whether longer distances are "
         "associated with higher shipping costs."),
        ("correlation_heatmap.png", "4.5 Correlation Heatmap",
         "The heatmap summarizes linear relationships between numeric "
         "variables. Correlation does not prove causation.")
    ]

    doc.add_heading("4. Visualizations and Interpretation", level=1)
    for filename, heading, explanation in chart_details:
        path = CHARTS / filename
        if path.exists():
            doc.add_heading(heading, level=2)
            doc.add_picture(str(path), width=Inches(5.7))
            doc.add_paragraph(explanation)

    doc.add_heading("5. Recommendations", level=1)
    for recommendation in [
        "Monitor actual delivery time against expected delivery time.",
        "Investigate delays by traffic, weather, destination, and warehouse.",
        "Compare shipping cost with distance and vehicle type.",
        "Track on-time delivery rate and average shipping cost regularly.",
        "Use a larger real-world dataset before making major operational changes."
    ]:
        doc.add_paragraph(recommendation, style="List Bullet")

    doc.add_heading("6. Limitations", level=1)
    doc.add_paragraph(
        "The dataset is small and illustrative. Its patterns may not represent "
        "all logistics operations. Correlations show associations, not causes."
    )

    doc.add_heading("7. Conclusion", level=1)
    doc.add_paragraph(
        "Exploratory data analysis and visualizations help summarize logistics "
        "performance and highlight areas for further investigation. More "
        "representative data would improve the reliability of future findings."
    )

    report = OUTPUT / "Week3_Logistics_Analysis_Report.docx"
    doc.save(report)

    print("\nWEEK 3 ANALYSIS COMPLETED")
    print("Total orders:", total_orders)
    print("Average shipping cost:", round(average_cost, 2))
    print("Average delivery days:", round(average_delivery, 2))
    if on_time_rate is not None:
        print("On-time delivery rate:", round(on_time_rate, 1), "%")
    print("Charts saved in:", CHARTS)
    print("Word report saved at:", report)


if __name__ == "__main__":
    main()
